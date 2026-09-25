"""Application packet generation from APPROVED facts only.

Every output unit carries provenance:
  fact_ids       -> approved facts the text was built from (verbatim fact values)
  profile_fields -> user-entered profile fields
  job_fields     -> sanitized job title/employer (never candidate claims)
  template       -> fixed boilerplate with no claims
`verify_provenance` is run on every packet and before approval. It fails if any unit is unsupported.
"""
import io
import re

from ..connectors.base import sanitize_short
from ..models import Fact, Job, Profile

MAX_ACHIEVEMENTS_PER_ROLE = 6
_INJECTION = re.compile(
    r"(ignore|disregard|forget)\b.{0,40}\b(instruction|previous|above|prompt)|system prompt|you are (an?|the) |"
    r"\bassistant\b|api[_ ]?key|password|<\s*script|https?://|\bexecute\b|\bsudo\b|\{\{|\}\}",
    re.IGNORECASE,
)


def looks_injected(s: str) -> bool:
    return bool(_INJECTION.search(s or "")) or len(s or "") > 150 or (s or "").count(".") > 2


def safe_job_fields(job: Job) -> tuple[str, str, list[str]]:
    """Title/employer as inserted into materials. Suspicious values are replaced, never passed through."""
    flags = []
    title, employer = sanitize_short(job.title, 150), sanitize_short(job.employer, 120)
    if looks_injected(title):
        flags.append("job title contains instruction-like or unusual text; replaced with a neutral phrase")
        title = "the advertised position"
    if looks_injected(employer):
        flags.append("employer name contains instruction-like or unusual text; replaced with a neutral phrase")
        employer = "your organization"
    return title, employer, flags


def _period(r: Fact) -> str:
    s, e = r.data.get("start") or "?", r.data.get("end") or "?"
    return f"{s} – {'Present' if e == 'present' else e}"


def _relevance(text: str, skills: list[str]) -> int:
    low = text.lower()
    return sum(1 for s in skills if s and re.search(rf"(?<![a-z0-9]){re.escape(s)}(?![a-z0-9])", low))


def build_resume_lines(profile: Profile, facts: list[Fact], matched_skills: list[str]) -> list[dict]:
    approved = [f for f in facts if f.status == "APPROVED"]
    L: list[dict] = []

    def add(text: str, section: str, **prov) -> None:
        L.append({"text": text, "section": section, "fact_ids": prov.get("fact_ids", []),
                  "profile_fields": prov.get("profile_fields", []), "template": prov.get("template", False)})

    add(f"{profile.first_name} {profile.last_name}", "header", profile_fields=["first_name", "last_name"])
    contact = [("email", profile.email), ("phone", profile.phone), ("linkedin_url", profile.linkedin_url),
               ("location", profile.location)]
    present = [(k, v) for k, v in contact if v]
    if present:
        add(" · ".join(v for _, v in present), "header", profile_fields=[k for k, _ in present])
    summary = next((f for f in approved if f.kind == "summary"), None)
    if summary:
        add(summary.data.get("text", ""), "summary", fact_ids=[summary.id])

    roles = sorted([f for f in approved if f.kind == "role"], key=lambda r: (r.data.get("start") or ""), reverse=True)
    if roles:
        add("Experience", "heading", template=True)
    for r in roles:
        emp = r.data.get("employer")
        add(f"{r.data.get('title', '')}{' — ' + emp if emp else ''} ({_period(r)})", "role", fact_ids=[r.id])
        kids = [f for f in approved if f.kind == "achievement" and f.parent_id == r.id]
        kids.sort(key=lambda k: _relevance(k.data.get("text", ""), matched_skills), reverse=True)
        for k in kids[:MAX_ACHIEVEMENTS_PER_ROLE]:
            add(k.data.get("text", ""), "achievement", fact_ids=[k.id])

    skills = [f for f in approved if f.kind == "skill"]
    if skills:
        ms = {m.lower() for m in matched_skills}
        skills.sort(key=lambda s: (0 if (s.data.get("name") or "").lower() in ms else 1, (s.data.get("name") or "").lower()))
        add("Skills", "heading", template=True)
        add(", ".join(s.data.get("name", "") for s in skills), "skills", fact_ids=[s.id for s in skills])

    edu = [f for f in approved if f.kind == "education"]
    if edu:
        add("Education", "heading", template=True)
        for e in edu:
            year = e.data.get("year")
            text_so_far = f"{e.data.get('degree') or ''} {e.data.get('institution') or ''}"
            bits = [e.data.get("degree"), e.data.get("institution"), year if year and str(year) not in text_so_far else None]
            add(" — ".join(str(b) for b in bits if b), "education", fact_ids=[e.id])
    certs = [f for f in approved if f.kind == "certification"]
    if certs:
        add("Certifications", "heading", template=True)
        for c in certs:
            add(c.data.get("name", ""), "certification", fact_ids=[c.id])
    projects = [f for f in approved if f.kind == "project"]
    if projects:
        add("Projects", "heading", template=True)
        for p in projects:
            add(p.data.get("text", ""), "project", fact_ids=[p.id])
    standalone = [f for f in approved if f.kind == "achievement" and f.parent_id is None]
    if standalone:
        add("Achievements", "heading", template=True)
        for a in standalone:
            add(a.data.get("text", ""), "achievement", fact_ids=[a.id])
    return L


def build_cover_letter(profile: Profile, job: Job, facts: list[Fact], matched_skill_fact_ids: dict) -> tuple[list[dict], list[str]]:
    approved = {f.id: f for f in facts if f.status == "APPROVED"}
    title, employer, flags = safe_job_fields(job)
    U: list[dict] = []

    def add(text: str, **prov) -> None:
        U.append({"text": text, "fact_ids": prov.get("fact_ids", []), "profile_fields": prov.get("profile_fields", []),
                  "job_fields": prov.get("job_fields", []), "template": prov.get("template", False)})

    add(f"Dear Hiring Team at {employer},", job_fields=["employer"])
    add(f"I am applying for the {title} position.", job_fields=["title"])
    roles = sorted([f for f in approved.values() if f.kind == "role"], key=lambda r: r.data.get("start") or "", reverse=True)
    if roles:
        r = roles[0]
        verb = "I currently work" if r.data.get("end") == "present" else "Most recently, I worked"
        emp = r.data.get("employer")
        add(f"{verb} as {r.data.get('title')}{' at ' + emp if emp else ''}.", fact_ids=[r.id])
    skill_names = sorted(matched_skill_fact_ids)
    achievements = [f for f in approved.values() if f.kind == "achievement"]
    achievements.sort(key=lambda a: _relevance(a.data.get("text", ""), skill_names), reverse=True)
    top = [a for a in achievements[:3] if a.data.get("text")]
    if top:
        add("Relevant highlights from my experience:", template=True)
        for a in top:
            add(f"• {a.data['text'].rstrip('.')}.", fact_ids=[a.id])
    if skill_names:
        ids = sorted({i for s in skill_names for i in matched_skill_fact_ids[s] if i in approved})
        names = [approved[i].data.get("name") for i in ids]
        if names:
            add(f"My experience includes {', '.join(names)}, which this role calls for.", fact_ids=ids)
    add("Thank you for your time and consideration.", template=True)
    add("Sincerely,", template=True)
    add(f"{profile.first_name} {profile.last_name}", profile_fields=["first_name", "last_name"])
    return U, flags


def verify_provenance(units: list[dict], approved_fact_ids: set[int]) -> list[str]:
    """Returns violations. Empty list = every unit is backed by approved facts, profile fields, job fields or template."""
    problems = []
    for i, u in enumerate(units):
        ids = u.get("fact_ids") or []
        if not (ids or u.get("profile_fields") or u.get("job_fields") or u.get("template")):
            problems.append(f"unit {i} has no provenance: {u.get('text', '')[:60]!r}")
        bad = [x for x in ids if x not in approved_fact_ids]
        if bad:
            problems.append(f"unit {i} cites unapproved facts {bad}")
    return problems


def render_docx(lines: list[dict]) -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for ln in lines:
        sec, text = ln["section"], ln["text"]
        if sec == "header" and ln is lines[0]:
            d.add_heading(text, level=0)
        elif sec == "heading":
            d.add_heading(text, level=2)
        elif sec in ("achievement", "project"):
            d.add_paragraph(text, style="List Bullet")
        elif sec == "role":
            p = d.add_paragraph()
            p.add_run(text).bold = True
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_cover_letter_text(units: list[dict]) -> str:
    out, prev_bullet = [], False
    for u in units:
        t = u["text"]
        is_bullet = t.startswith("• ")
        out.append(("\n" if is_bullet and prev_bullet else "\n\n") + t if out else t)
        prev_bullet = is_bullet
    return "".join(out).strip() + "\n"
