"""Isolated document parser. Runs as a separate process with resource limits and no secrets in env.

Usage (internal): python -m app.ingestion.parse_worker <pdf|docx>   (bytes on stdin, JSON on stdout)
Keep imports minimal: this module must not import app settings or crypto.
"""
import io
import json
import sys
import zipfile

MAX_ZIP_ENTRIES = 2000
MAX_ZIP_UNCOMPRESSED = 60 * 1024 * 1024
MAX_TEXT_CHARS = 400_000


def parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("encrypted PDFs are not supported")
    if len(reader.pages) > 30:
        raise ValueError("resume has more than 30 pages")
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def _check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = zf.infolist()
        if len(infos) > MAX_ZIP_ENTRIES:
            raise ValueError("too many archive entries")
        if sum(i.file_size for i in infos) > MAX_ZIP_UNCOMPRESSED:
            raise ValueError("archive expands too large")


def parse_docx(data: bytes) -> str:
    _check_zip(data)
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return "\n".join(lines)


def _limit_resources(mem_mb: int, cpu_s: int) -> None:
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (cpu_s, cpu_s))
        try:
            resource.setrlimit(resource.RLIMIT_AS, (mem_mb * 1024 * 1024, mem_mb * 1024 * 1024))
        except (ValueError, OSError):
            pass  # RLIMIT_AS is not enforceable on macOS; the wall-clock timeout still applies
    except ImportError:
        pass


def main() -> int:
    kind = sys.argv[1]
    mem_mb = int(sys.argv[2]) if len(sys.argv) > 2 else 512
    cpu_s = int(sys.argv[3]) if len(sys.argv) > 3 else 20
    _limit_resources(mem_mb, cpu_s)
    data = sys.stdin.buffer.read()
    try:
        text = parse_pdf(data) if kind == "pdf" else parse_docx(data)
        text = text.replace("\x00", "")[:MAX_TEXT_CHARS]
        sys.stdout.write(json.dumps({"ok": True, "text": text}))
    except Exception as e:  # noqa: BLE001 - report any parser failure as data
        sys.stdout.write(json.dumps({"ok": False, "error": f"{type(e).__name__}: {str(e)[:300]}"}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
